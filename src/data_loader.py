from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


@dataclass(frozen=True)
class DataValidation:
    row_count: int
    trailing_ma30_rows: int
    trailing_ma30_matches: int
    max_ma30_error: float
    formula_rows_checked: int
    formula_rows_matching: int
    formula_check_available: bool
    note_cells: list[tuple[str, str]]


CONDITION_LABELS = {
    "normal_caving": "正常放煤",
    "minor_caving": "少量放煤",
    "excessive_caving": "过量放煤",
}


def load_radiation_data(path: str | Path) -> pd.DataFrame:
    """Load the teacher-provided count-rate workbook into normalized columns."""
    path = Path(path)
    raw = pd.read_excel(path, sheet_name="Sheet1", header=None, engine="openpyxl")
    df = raw.iloc[:, :3].copy()
    df.columns = ["sample_index", "count_rate", "ma30_count_rate_source"]
    df = df[pd.to_numeric(df["sample_index"], errors="coerce").notna()]
    df = df[pd.to_numeric(df["count_rate"], errors="coerce").notna()]
    df["sample_index"] = df["sample_index"].astype(float).astype(int)
    df["count_rate"] = df["count_rate"].astype(float)
    df["ma30_count_rate_source"] = pd.to_numeric(
        df["ma30_count_rate_source"], errors="coerce"
    )
    df = df.reset_index(drop=True)
    df["row_number"] = np.arange(1, len(df) + 1)
    return df[["row_number", "sample_index", "count_rate", "ma30_count_rate_source"]]


def parse_caving_condition(file_name: str) -> tuple[str, str]:
    """Parse caving condition from an English-named caving test Excel file."""
    for condition, chinese_label in CONDITION_LABELS.items():
        if condition in file_name.lower():
            return condition, chinese_label
    return "unknown", "unknown"


def load_caving_condition_workbook(path: str | Path, sample_interval_s: float = 0.1) -> pd.DataFrame:
    """Load one single-column caving test workbook as a count-rate time series."""
    path = Path(path)
    raw = pd.read_excel(path, sheet_name="Sheet1", header=None, engine="openpyxl")
    values = pd.to_numeric(raw.iloc[:, 0], errors="coerce").dropna().astype(float).reset_index(drop=True)
    condition, condition_label = parse_caving_condition(path.name)
    out = pd.DataFrame(
        {
            "row_number": np.arange(1, len(values) + 1),
            "sample_index": np.arange(len(values)),
            "time_s": np.arange(len(values), dtype=float) * float(sample_interval_s),
            "count_rate": values,
            "condition": condition,
            "condition_label": condition_label,
            "source_file": path.name,
        }
    )
    return out


def load_caving_condition_data(
    directory: str | Path,
    sample_interval_s: float = 0.1,
) -> pd.DataFrame:
    """Load all caving-condition Excel files from a directory."""
    directory = Path(directory)
    frames = [
        load_caving_condition_workbook(path, sample_interval_s=sample_interval_s)
        for path in sorted(directory.glob("*.xlsx"), key=lambda item: item.name)
    ]
    if not frames:
        raise FileNotFoundError(f"No .xlsx files found in {directory}")
    return pd.concat(frames, ignore_index=True)


def validate_workbook(path: str | Path) -> DataValidation:
    """Verify that workbook column C is a trailing 30-point moving average."""
    path = Path(path)
    wb_values = load_workbook(path, data_only=True, read_only=True)
    ws_values = wb_values["Sheet1"]
    wb_formula = load_workbook(path, data_only=False, read_only=True)
    ws_formula = wb_formula["Sheet1"]

    note_cells: list[tuple[str, str]] = []
    for row in ws_values.iter_rows():
        for cell in row:
            if cell.value is not None and cell.column > 3:
                note_cells.append((cell.coordinate, str(cell.value)))

    count_values: list[float] = []
    source_ma: list[tuple[int, float]] = []
    formula_rows_checked = 0
    formula_rows_matching = 0
    for row_idx in range(1, ws_values.max_row + 1):
        count = ws_values.cell(row_idx, 2).value
        ma_value = ws_values.cell(row_idx, 3).value
        formula = ws_formula.cell(row_idx, 3).value
        if isinstance(count, (int, float)):
            count_values.append(float(count))
        if isinstance(ma_value, (int, float)):
            source_ma.append((row_idx, float(ma_value)))
        if isinstance(formula, str) and formula.startswith("="):
            formula_rows_checked += 1
            expected_formula = f"=SUM(B{row_idx - 29}:B{row_idx})/30"
            expected_average = f"=AVERAGE(B{row_idx - 29}:B{row_idx})"
            normalized = formula.replace(" ", "").upper()
            if row_idx >= 30 and normalized in {expected_formula.upper(), expected_average.upper()}:
                formula_rows_matching += 1

    matches = 0
    max_error = 0.0
    for row_idx, ma_value in source_ma:
        if row_idx < 30:
            continue
        expected = float(np.mean(count_values[row_idx - 30 : row_idx]))
        error = abs(expected - ma_value)
        max_error = max(max_error, error)
        if error < 1e-9:
            matches += 1

    return DataValidation(
        row_count=len(count_values),
        trailing_ma30_rows=len(source_ma),
        trailing_ma30_matches=matches,
        max_ma30_error=max_error,
        formula_rows_checked=formula_rows_checked,
        formula_rows_matching=formula_rows_matching,
        formula_check_available=formula_rows_checked > 0,
        note_cells=note_cells,
    )
