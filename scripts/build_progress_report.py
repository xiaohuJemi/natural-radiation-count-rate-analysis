from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "progress_report.docx"
SUMMARY = ROOT / "outputs" / "analysis_summary.json"
FIG1 = ROOT / "outputs" / "fig01_time_series.png"
FIG4 = ROOT / "outputs" / "fig04_changepoints.png"
FIG6 = ROOT / "outputs" / "fig06_method_comparison.png"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(9)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.color.rgb = RGBColor(31, 78, 121)


def add_para(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r2 = p.add_run(text[len(bold_prefix) :])
        runs = [r1, r2]
    else:
        runs = [p.add_run(text)]
    for run in runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)
        set_cell_shading(table.rows[0].cells[i], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    if widths:
        for row in table.rows:
            for idx, width_cm in enumerate(widths):
                row.cells[idx].width = Cm(width_cm)
    doc.add_paragraph()


def add_picture_with_caption(doc: Document, path: Path, caption: str) -> None:
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.1))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(89, 89, 89)


def apply_base_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)


def build_doc() -> None:
    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))
    doc = Document()
    apply_base_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("综放工作面自然射线计数率探索性时序分析\n当前进度汇报")
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("进度汇报")
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(89, 89, 89)

    add_heading(doc, "一、研究背景与当前定位", 1)
    add_para(
        doc,
        "本研究围绕综采/综放工作面煤矸识别问题展开，前期重点尝试利用自然射线计数率时间序列，结合滑动统计、阈值分析、变点检测和无监督状态划分，探索煤流、夹矸混入和高计数异常阶段的可能识别方法。",
    )
    add_para(
        doc,
        "需要说明的是：根据近期专业审阅意见和对项目本身的复核，目前工作应定位为“探索性数据分析和预实验”，尚不能直接表述为已经完成可发表级别的煤矸智能识别方法。主要原因是当前只有一组无标签计数率序列，缺少现场标签、传感器参数和多组数据验证。",
    )

    add_heading(doc, "二、已完成工作", 1)
    add_bullets(
        doc,
        [
            "整理并阅读 references 文件夹中的 4 篇煤矸识别与自然射线相关文献，形成文献研读报告。",
            "完成 datas/data.xlsx 的读取和数据结构校验，确认核心数据为一条自然射线计数率序列。",
            "用 Python 建立了可复现实验流程，包括数据读取、特征构造、阈值分析、变点检测、GMM 状态划分和异常点筛选。",
            "生成 7 张图件，包括原始计数率曲线、分布图、滑动统计图、变点检测图、GMM 状态图、方法对比图和概念示意图。",
            "完成论文初稿、算法报告、投稿建议和专家审阅后的下一阶段 planning。",
        ],
    )

    add_heading(doc, "三、现有数据基本情况", 1)
    data = summary["data"]
    validation = summary["workbook_validation"]
    add_table(
        doc,
        ["项目", "当前结果", "说明"],
        [
            ["样本数", f"{data['rows']} 行", "当前仅有一组序列"],
            ["采样序号范围", f"{data['sample_index_min']} - {data['sample_index_max']}", "序号步长为 2，原因待确认"],
            ["原始计数率范围", f"{data['count_rate_min']:.0f} - {data['count_rate_max']:.0f}", "第 2 列为核心原始数据"],
            ["原始计数率均值", f"{data['count_rate_mean']:.2f}", "仅为描述性统计"],
            ["30 点移动平均范围", f"{data['ma30_min']:.2f} - {data['ma30_max']:.2f}", "由第 2 列计算得到"],
            ["第 3 列校验", f"{validation['trailing_ma30_matches']}/{validation['trailing_ma30_rows']} 完全匹配", "第 3 列不是独立传感器数据"],
        ],
        widths=[3.2, 4.2, 7.2],
    )

    add_picture_with_caption(doc, FIG1, "图1 现有自然射线计数率序列及 30 点移动平均")

    add_heading(doc, "四、初步算法分析结果", 1)
    add_para(
        doc,
        "当前代码将原始计数率构造成多尺度移动平均、滑动标准差、差分、斜率和分位数等特征，并进行了阈值法、动态规划变点检测、GMM 聚类和 Isolation Forest 高计数异常筛选。结果显示序列中部存在明显高计数阶段，但由于缺少人工标签，暂不能证明该阶段一定对应顶板岩石或夹矸混入。",
    )
    add_table(
        doc,
        ["方法", "当前结果", "可信解释"],
        [
            ["阈值法", "基于早期低计数段得到多个阈值", "只能作为探索性分层；当前阈值过于敏感，需要重新设计"],
            ["变点检测", "将序列分成 8 段", "说明序列存在阶段变化；但 8 段受参数上限影响，需要修正 BIC 和敏感性分析"],
            ["GMM", "4 类状态，轮廓系数 0.245", "聚类结构较弱，不宜作为强识别结论"],
            ["Isolation Forest", "筛出 33 个高计数异常点", "只能称为高计数异常，不能直接称为顶板岩石混入"],
        ],
        widths=[3.2, 5.2, 6.2],
    )
    add_picture_with_caption(doc, FIG4, "图2 当前变点检测结果：仅表示计数率阶段边界的探索性结果")

    add_heading(doc, "五、专家审阅后的主要问题", 1)
    add_bullets(
        doc,
        [
            "数据层面：当前只有一组无标签序列，缺少训练/测试划分、交叉验证和泛化评估。",
            "标签层面：缺少放煤窗口动作日志、人工观察记录、视频标注或采样化验结果，因此无法验证“高计数率对应煤矸状态变化”的结论。",
            "传感器层面：探测器型号、采样周期、测量窗口、安装位置、距煤流距离、本底计数率等信息缺失。",
            "方法层面：BIC 公式和分段上限需要修正；GMM 组件数不应固定为 4；阈值基线选择需要统计依据。",
            "论文表述层面：当前应降级为探索性分析，不应直接称为“智能识别方法”或“可投稿成果”。",
        ],
    )

    add_heading(doc, "六、下一步工作计划", 1)
    add_table(
        doc,
        ["阶段", "工作内容", "目标"],
        [
            ["第 1 阶段", "修复代码统计问题，补充依赖声明、测试和敏感性分析", "使当前分析流程更严谨、可复现"],
            ["第 2 阶段", "向之前大创项目联系过的老师补充询问原始数据、传感器参数和现场记录", "补齐数据来源和实验条件"],
            ["第 3 阶段", "争取获得 3-5 组独立放煤过程数据及至少一种标签或验证依据", "从探索性分析升级为可验证研究"],
            ["第 4 阶段", "根据是否获得标签决定论文定位", "有标签则写识别方法；无标签则写预实验/探索性报告"],
        ],
        widths=[2.5, 7.0, 5.1],
    )

    add_heading(doc, "七、需要补充获取的数据与信息", 1)
    add_para(doc, "考虑到提供数据的老师并非本人导师，而是此前大创项目中联系过的老师，沟通时应尽量礼貌、轻量，说明是为了确认数据背景和完善本科阶段研究，不宜一次性提出过重要求。")
    add_bullets(
        doc,
        [
            "该 Excel 数据的采集背景：矿井/工作面名称、采集时间、是否为现场数据或实验数据。",
            "第 1 列采样序号的含义：为什么为 1、3、5 递增，原始采样周期是多少。",
            "探测器相关信息：型号、类型、安装位置、距煤流距离、测量时间窗口、采样率、本底计数率。",
            "是否有该次采集对应的放煤窗口开闭记录、人工观察记录、视频或文字说明。",
            "该数据中哪些阶段可能对应煤流为主、夹矸混入、顶板岩石混入或放煤结束。",
            "是否还有其他几组类似的自然射线计数率原始数据。",
            "如果方便，是否能提供当时项目或实验记录中对数据列含义、采集条件的说明。",
        ],
    )

    add_heading(doc, "八、阶段性结论", 1)
    add_para(
        doc,
        "目前项目已经完成文献梳理、数据初步分析和可复现代码框架，具备继续推进的基础。但从论文研究角度看，当前最大短板不是代码或图表，而是数据证据链不足。下一步应优先补充数据来源、传感器参数和 ground truth。如果无法获得这些信息，项目更适合定位为本科阶段的探索性分析或预实验报告；如果能够获得多组有标签数据，则可以继续向煤矸识别方法论文推进。",
    )

    section = doc.add_section(WD_SECTION.CONTINUOUS)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("煤矸自然射线识别研究进度汇报")
    fr.font.name = "Microsoft YaHei"
    fr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    fr.font.size = Pt(9)
    fr.font.color.rgb = RGBColor(128, 128, 128)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_doc()

