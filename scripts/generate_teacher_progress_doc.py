import csv
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageOps


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "progress_report_to_teacher_detailed.docx"
FIG_DIR = ROOT / "reports" / "progress_report_figures"


FONT_CN = "Microsoft YaHei"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 78, 121)
GRAY_FILL = "F2F4F7"
LIGHT_BLUE_FILL = "EAF2F8"
NOTE_FILL = "FFF7E6"


def read_csv_rows(path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def fmt_float(value, digits=2, empty=""):
    if value in (None, ""):
        return empty
    return f"{float(value):.{digits}f}"


def fmt_pct(value, digits=1, empty=""):
    if value in (None, ""):
        return empty
    return f"{float(value) * 100:.{digits}f}%"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, size=9.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = FONT_CN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    run.font.size = Pt(size)


def style_paragraph(paragraph, size=10.5, bold=False, color=None, after=5, before=0):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.12
    for run in paragraph.runs:
        run.font.name = FONT_CN
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        run.font.size = Pt(size)
        run.bold = bold
        if color:
            run.font.color.rgb = color


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles[f"Heading {level}"]
    run = paragraph.add_run(text)
    size = {1: 15, 2: 12.5, 3: 11.5}.get(level, 11)
    color = BLUE if level < 3 else DARK_BLUE
    run.font.name = FONT_CN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = color
    paragraph.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    paragraph.paragraph_format.space_after = Pt(4)
    return paragraph


def add_body(doc, text, after=5):
    paragraph = doc.add_paragraph(text)
    style_paragraph(paragraph, after=after)
    return paragraph


def add_callout(doc, title, text, fill=LIGHT_BLUE_FILL):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(title + "：")
    r1.bold = True
    r1.font.name = FONT_CN
    r1._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    r1.font.size = Pt(10.5)
    r2 = p.add_run(text)
    r2.font.name = FONT_CN
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    r2.font.size = Pt(10.5)
    doc.add_paragraph()


def add_bullet_list(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)
        style_paragraph(paragraph, size=10.3, after=3)


def add_numbered_list(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.add_run(item)
        style_paragraph(paragraph, size=10.3, after=3)


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, header, bold=True, size=font_size)
        set_cell_shading(cell, GRAY_FILL)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size)
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.TOP

    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph()
    return table


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.name = FONT_CN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(89, 89, 89)
    paragraph.paragraph_format.space_after = Pt(6)


def add_figure(doc, path, caption, width_cm=15.5):
    if not path.exists():
        add_callout(doc, "图件缺失", f"未找到图件：{path}", fill=NOTE_FILL)
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    add_caption(doc, caption)


def make_vertical_composite(paths, output_path, labels):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    images = []
    target_width = 2200
    for path, label in zip(paths, labels):
        with Image.open(path) as raw:
            image = raw.convert("RGB")
            ratio = target_width / image.width
            resized = image.resize((target_width, int(image.height * ratio)), Image.LANCZOS)
            canvas = ImageOps.expand(resized, border=(0, 90, 0, 18), fill="white")
            images.append((canvas, label))

    font_margin = 34
    total_height = sum(img.height for img, _ in images)
    composite = Image.new("RGB", (target_width, total_height), "white")
    y = 0
    for img, label in images:
        composite.paste(img, (0, y))
        # Use simple raster text only as panel label; doc caption gives full interpretation.
        from PIL import ImageDraw

        draw = ImageDraw.Draw(composite)
        draw.text((font_margin, y + 22), label, fill=(31, 78, 121))
        y += img.height
    composite.save(output_path, quality=95)
    return output_path


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.55)
    section.right_margin = Cm(2.55)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    for style_name in ["Normal", "List Bullet", "List Number"]:
        style = doc.styles[style_name]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(10.5)

    for level in [1, 2, 3]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.bold = True


def add_title_block(doc):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("当前研究进度详细汇报与请教事项")
    run.font.name = FONT_CN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = DARK_BLUE

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("基于自然射线计数率时间序列的综放煤矸识别探索")
    run.font.name = FONT_CN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    run.font.size = Pt(11.5)
    run.font.color.rgb = RGBColor(89, 89, 89)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("汇报人：________（本科生）    日期：____年__月__日")
    run.font.name = FONT_CN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    run.font.size = Pt(10.2)


def build_composites():
    figures = {
        "mechanism": ROOT / "outputs" / "concept01_radiation_mechanism_diagram.png",
        "early_series": ROOT / "outputs" / "fig01_time_series.png",
        "early_compare": ROOT / "outputs" / "fig06_method_comparison.png",
        "background_range": ROOT / "outputs" / "new_data" / "fig_background_range_by_strategy.png",
        "component_strategy": ROOT / "outputs" / "new_data" / "fig_radiation_component_by_strategy.png",
        "dynamic_crossing": ROOT / "outputs" / "new_data" / "fig_dynamic_crossing_by_strategy.png",
        "poisson_summary": ROOT / "outputs" / "new_data" / "fig_poisson_excess_summary.png",
    }
    figures["background_overlay"] = make_vertical_composite(
        [
            ROOT / "outputs" / "new_data" / "fig_background_overlay_normal_caving.png",
            ROOT / "outputs" / "new_data" / "fig_background_overlay_minor_caving.png",
            ROOT / "outputs" / "new_data" / "fig_background_overlay_excessive_caving.png",
        ],
        FIG_DIR / "fig_report_background_overlay.png",
        ["A  Normal caving", "B  Minor caving", "C  Excessive caving"],
    )
    figures["poisson_fluctuation"] = make_vertical_composite(
        [
            ROOT / "outputs" / "new_data" / "fig_poisson_fluctuation_normal_caving.png",
            ROOT / "outputs" / "new_data" / "fig_poisson_fluctuation_minor_caving.png",
            ROOT / "outputs" / "new_data" / "fig_poisson_fluctuation_excessive_caving.png",
        ],
        FIG_DIR / "fig_report_poisson_fluctuation.png",
        ["A  Normal caving", "B  Minor caving", "C  Excessive caving"],
    )
    figures["stage_segments"] = make_vertical_composite(
        [
            ROOT / "outputs" / "new_data" / "fig_stage_segments_normal_caving.png",
            ROOT / "outputs" / "new_data" / "fig_stage_segments_minor_caving.png",
            ROOT / "outputs" / "new_data" / "fig_stage_segments_excessive_caving.png",
        ],
        FIG_DIR / "fig_report_stage_segments.png",
        ["A  Normal caving", "B  Minor caving", "C  Excessive caving"],
    )
    figures["csv_series"] = make_vertical_composite(
        [
            ROOT / "outputs" / "new_data" / "fig_csv_minute_trend.png",
            ROOT / "outputs" / "new_data" / "fig_csv_background_components.png",
            ROOT / "outputs" / "new_data" / "fig_csv_high_count_events.png",
        ],
        FIG_DIR / "fig_report_csv_series.png",
        ["A  Minute-level trend", "B  Background and component", "C  High-count events"],
    )
    return figures


def build_document():
    figures = build_composites()
    caving_rows = read_csv_rows(ROOT / "outputs" / "new_data" / "caving_condition_summary.csv")
    poisson_rows = [
        row
        for row in read_csv_rows(ROOT / "outputs" / "new_data" / "caving_poisson_summary.csv")
        if row["confidence_level"] == "0.99"
    ]
    csv_summary = read_csv_rows(ROOT / "outputs" / "new_data" / "csv_long_series_summary.csv")[0]

    doc = Document()
    configure_document(doc)
    add_title_block(doc)

    add_callout(
        doc,
        "汇报说明",
        "老师您好！我是之前参与大创项目时与您有过联系的本科生。近期我围绕自然射线计数率在综放工作面煤矸识别中的应用做了阶段性整理。本汇报不是最终论文结论，而是把目前能够从时间序列中得到的结果、已有疑问和下一步计划集中呈现，希望请您从现场理解、数据使用和研究路线三个方面给予指导。",
    )

    add_heading(doc, "一、研究目标与当前边界", 1)
    add_body(
        doc,
        "当前研究只分析自然射线计数率随时间变化的规律。根据老师此前说明，早期 data 数据、新增 Excel 数据和新增 CSV 数据来源不同，三者没有绝对对应关系；文件名中的编号暂不作为物理参数解释；所有序列按相邻两行间隔 0.1 s 处理。因此，本阶段不尝试建立跨数据源的统一绝对阈值，而是分别讨论各数据内部的时间序列特征。",
    )
    add_bullet_list(
        doc,
        [
            "研究对象：自然射线计数率时间序列，而不是设备电压、支架编号或其他辅助字段。",
            "核心问题：如何从计数率中分离本底影响与煤矸流附加辐射影响，并判断高计数是否超过放射性统计涨落。",
            "当前定位：无强人工标签条件下的机理约束时序分析，属于本科阶段较稳妥的研究路线。",
            "主要风险：缺少空载本底、人工标注放煤阶段、支架动作时序和煤矸比例记录，因此结果必须谨慎表述。",
        ]
    )

    add_figure(
        doc,
        figures["mechanism"],
        "图1  自然射线煤矸识别机理示意。煤、矸石和顶板岩石因天然放射性核素含量差异造成计数率差异；实际识别时还会受到距离、屏蔽、探测体积、环境本底和统计涨落影响。",
    )

    add_heading(doc, "二、已完成的主要工作", 1)
    add_table(
        doc,
        ["工作模块", "具体完成内容", "阶段性价值"],
        [
            [
                "文献与机理整理",
                "梳理自然 γ 射线识别、低水平计数统计涨落、混矸率与辐射强度关系、放煤阶段时序规律等内容，并重点参考韦明辉论文中的实验和阈值思路。",
                "明确不能把单点高值直接等同于矸石，应结合本底、持续时间和置信区间判断。",
            ],
            [
                "数据边界梳理",
                "将早期 data、新增 Excel、新增 CSV 分开处理；统一保留时间序列计数率，按 0.1 s 间隔重建时间轴。",
                "避免把不同来源数据强行拼接或使用同一绝对阈值解释。",
            ],
            [
                "本底拆分",
                "建立 B(t)+R(t) 的分析框架：B(t) 表示本底或缓慢环境变化，R(t) 表示相对本底的附加辐射贡献。",
                "回应“工作面推进和环境变化会影响本底”的问题。",
            ],
            [
                "泊松涨落分析",
                "按放射性计数服从泊松统计的特点，计算 95% 和 99% 置信上限，判断计数率升高是否超出随机涨落范围。",
                "比简单 99% 分位数筛查更有物理依据。",
            ],
            [
                "阶段分割与可视化",
                "生成原始序列、滑动均值、本底曲线、附加辐射分量、泊松置信区间、阶段分割和长序列趋势图。",
                "形成论文图件和向老师汇报的直观证据。",
            ],
        ],
        widths=[3.0, 7.2, 5.2],
    )

    add_heading(doc, "三、早期 data 数据的整体分析", 1)
    add_body(
        doc,
        "早期 data 数据共有 1407 行，核心是原始计数率和 30 点滑动平均。该数据的计数率约在 40-166 之间，均值约 100.25，30 点滑动平均约在 60.37-144.23 之间，存在明显阶段变化。该部分主要用于验证基础流程：平滑、分布分析、变点检测、聚类状态识别和方法对比。",
    )
    add_figure(
        doc,
        figures["early_series"],
        "图2  早期 data 数据计数率时间序列。原始计数率存在明显波动，30 点滑动均值能压制单点随机起伏并突出阶段性变化。",
    )
    add_figure(
        doc,
        figures["early_compare"],
        "图3  阈值法、变点检测和状态聚类结果对比。该图说明单一阈值能识别高计数区间，但变点与聚类更适合描述时序阶段变化。",
    )

    add_heading(doc, "四、新增 Excel 工况数据分析", 1)
    add_body(
        doc,
        "新增 Excel 数据包括“正常放煤、少量放煤、过量放煤”三组时间序列。这里将文件名中的工况表述暂作为弱标签，仅用于对比不同序列内部特征，不把它当作已经严格标注的现场真值。三组数据的平均计数率接近，但高值持续时间、本底范围和泊松超限持续时间存在差异。",
    )
    add_table(
        doc,
        ["工况", "时长/s", "均值/cps", "最大值/cps", "ma10 最大值", "85.4 cps 最长越限/s", "滚动本底范围/cps"],
        [
            [
                row["condition_label"],
                fmt_float(row["duration_s"], 1),
                fmt_float(row["count_rate_mean"], 2),
                fmt_float(row["count_rate_max"], 1),
                fmt_float(row["ma10_max"], 1),
                fmt_float(row["fixed_longest_run_duration_s"], 1),
                fmt_float(row["background_range"], 2),
            ]
            for row in caving_rows
        ],
        widths=[2.2, 1.8, 2.0, 2.0, 2.0, 2.8, 2.4],
        font_size=8.5,
    )
    add_body(
        doc,
        "从表中看，三组 Excel 数据的均值差异并不大，说明若只比较平均计数率，很难区分放煤状态；更有信息量的是高计数持续时间、超过本底的附加分量、以及是否连续超过泊松置信上限。",
    )

    add_heading(doc, "五、本底值与煤矸附加贡献拆分", 1)
    add_body(
        doc,
        "本阶段采用 B(t)+R(t) 的分解思想。观测计数率或滑动计数率记为 X(t)，其中 B(t) 表示本底项，可能随工作面推进、探测器周围岩层、设备位置和环境变化缓慢漂移；R(t)=X(t)-B(t) 表示相对本底的附加辐射贡献。由于缺少空载本底，当前只能估计相对本底，不能断言 R(t) 完全来自矸石。",
    )
    add_bullet_list(
        doc,
        [
            "初始稳定段本底：取开头若干秒作为基准，优点是物理含义直观；缺点是若工作面环境随时间变化，本底会被固定住。",
            "滚动低分位本底：用局部低计数包络近似环境本底，优点是能跟随缓慢变化；缺点是可能把真实低煤矸贡献也吸收到本底中。",
            "指数平滑本底：强调平滑跟踪，适合描述缓慢漂移；缺点是参数选择会影响附加分量大小。",
        ]
    )
    add_figure(
        doc,
        figures["background_range"],
        "图4  不同本底估计策略下的本底范围。该图用于说明本底模型选择会直接影响后续附加辐射分量和阈值判别。",
    )
    add_figure(
        doc,
        figures["component_strategy"],
        "图5  不同本底估计策略下的附加辐射分量比较。若本底估计过于激进，部分煤矸影响会被吸收到 B(t) 中；若过于保守，则环境漂移会被误认为 R(t)。",
    )
    add_figure(
        doc,
        figures["background_overlay"],
        "图6  三种放煤工况下计数率、本底和附加分量叠加图。该图直观展示了 B(t) 与 R(t) 的拆分逻辑，也是后续论文中解释“本底影响”和“煤矸影响”的核心图件。",
    )

    add_heading(doc, "六、泊松涨落与置信区间分析", 1)
    add_body(
        doc,
        "自然射线计数属于随机计数过程。若在某一时间窗口内本底期望计数为 λ，则观测计数可近似服从 Poisson(λ)。因此，高计数是否有意义，不应只看数值是否大，而应判断其是否超过给定置信水平下的本底统计涨落范围。本阶段把 10 个 0.1 s 采样点理解为约 1 s 的平滑窗口，对 ma10 计数率建立 95% 和 99% 上置信界。",
    )
    add_table(
        doc,
        ["工况", "本底均值/cps", "99% 上限均值/cps", "ma10 最大值", "最大超出/cps", "99% 超限比例", "最长连续超限/s"],
        [
            [
                row["condition_label"],
                fmt_float(row["background_mean"], 2),
                fmt_float(row["poisson_upper_mean"], 1),
                fmt_float(row["max_ma10_count_rate"], 1),
                fmt_float(row["max_excess_over_poisson"], 1),
                fmt_pct(row["crossing_ratio"], 1),
                fmt_float(row["longest_run_duration_s"], 1),
            ]
            for row in poisson_rows
        ],
        widths=[2.1, 2.3, 2.4, 2.0, 2.0, 2.2, 2.4],
        font_size=8.3,
    )
    add_body(
        doc,
        "当前结果显示，过量放煤序列在 99% 置信上限外的连续高计数持续时间更长，说明其高值段较难仅用统计涨落解释。正常放煤和少量放煤也存在短时超限，因此不能仅凭是否超限判断状态，还需要结合持续时间、超限面积和现场标签。",
    )
    add_figure(
        doc,
        figures["poisson_summary"],
        "图7  三种工况的泊松超限统计摘要。该图把超限比例、最长连续超限时间和最大超出幅度放在一起，便于比较统计波动之外的高计数贡献。",
    )
    add_figure(
        doc,
        figures["poisson_fluctuation"],
        "图8  三种工况下 ma10 计数率与泊松置信上限。连续越过 99% 上限的区间是判断真实高计数贡献的重要候选段。",
    )

    add_heading(doc, "七、阶段分割与放煤状态解释", 1)
    add_body(
        doc,
        "在没有强标签的情况下，阶段分割的目标不是直接给出“煤/矸/岩”的最终分类，而是找出时间序列内部的稳定段、升高段和高值持续段。后续若能获得放煤口动作时间、人工记录或煤矸比例，就可以把这些阶段与现场事件对应起来。",
    )
    add_figure(
        doc,
        figures["stage_segments"],
        "图9  三种放煤工况的阶段分割结果。该图用于识别从稳定波动到持续高计数的转换区间，为后续建立弱监督标签或响应时间指标提供基础。",
    )
    add_callout(
        doc,
        "谨慎解释",
        "当前阶段分割只能说明计数率时间序列出现了状态变化，不能单独证明某一段一定是矸石或顶板岩石混入。若要提升结论强度，需要现场同步标签或更多重复试验。",
        fill=NOTE_FILL,
    )

    add_heading(doc, "八、CSV 长序列数据分析", 1)
    add_body(
        doc,
        "新增 CSV 数据规模更大，共 76867 行。其时间戳范围为 2025-11-28 17:26:18 至 19:40:27，按时间戳计算跨度约 8049 s；计数率范围为 126-348，均值约 206.92，99% 分位数为 282。该数据与 Excel 数据来源不同，均值水平明显更高，因此只在 CSV 内部研究趋势、本底漂移和高计数事件，不与 Excel 的绝对阈值直接比较。",
    )
    add_table(
        doc,
        ["指标", "数值"],
        [
            ["样本数", csv_summary["rows"]],
            ["时间跨度/s", fmt_float(csv_summary["duration_s_by_timestamp"], 1)],
            ["计数率范围/cps", f"{fmt_float(csv_summary['count_rate_min'], 1)} - {fmt_float(csv_summary['count_rate_max'], 1)}"],
            ["均值/cps", fmt_float(csv_summary["count_rate_mean"], 2)],
            ["标准差/cps", fmt_float(csv_summary["count_rate_std"], 2)],
            ["中位数/cps", fmt_float(csv_summary["count_rate_median"], 1)],
            ["99% 分位数/cps", fmt_float(csv_summary["count_rate_q99"], 1)],
            ["超过 99% 分位数样本比例", fmt_pct(csv_summary["above_high_count_ratio"], 2)],
        ],
        widths=[5.0, 5.0],
    )
    add_figure(
        doc,
        figures["csv_series"],
        "图10  CSV 长序列趋势、本底拆分和高计数事件。该数据更适合研究较长时间尺度上的环境本底变化和局部高计数事件，但目前缺少支架动作和现场标签，暂不用于判断前序支架影响。",
    )

    add_heading(doc, "九、对三个关键研究问题的当前回答", 1)
    add_heading(doc, "1. 本底值影响与煤矸影响如何拆分", 2)
    add_body(
        doc,
        "当前可行做法是先估计随时间变化的本底 B(t)，再用相对分量 R(t)=X(t)-B(t) 表示高于本底的附加贡献。为了避免误判，R(t) 还要经过泊松置信区间检验：只有持续超过高置信上限的片段，才可作为“真实高计数贡献”的候选段。该方法能够把研究从简单阈值提高到“本底校正 + 统计显著性 + 持续时间”的组合判据。",
    )
    add_heading(doc, "2. 前序支架放煤对当前探测点的影响如何考虑", 2)
    add_body(
        doc,
        "目前只有单点计数率时间序列，缺少支架编号、放煤口动作时间、刮板输送机速度、煤矸流到达探测点的输送延迟等信息。因此该问题暂时不能可靠建模。当前报告中只把它作为重要局限说明，后续若能获得多支架同步动作记录，可按“前序支架事件时间 + 输送延迟 + 当前探测点响应”的思路建立滞后相关分析。",
    )
    add_heading(doc, "3. 工作面推进导致环境变化时如何处理本底", 2)
    add_body(
        doc,
        "工作面推进会改变探测器周围岩层、空间几何、屏蔽条件和环境放射性背景，因此固定本底可能不适用于长时间连续监测。当前采用滚动低分位和指数平滑作为动态本底的近似方法，但由于没有空载本底或推进位置记录，仍只能作为工程近似。更严格的做法是补充不同推进位置下的无煤流本底，建立按时间或推进位置更新的本底校正模型。",
    )

    add_heading(doc, "十、需要请老师指导的问题", 1)
    add_numbered_list(
        doc,
        [
            "新增 Excel 文件名中的“正常放煤、少量放煤、过量放煤”能否作为弱标签使用？这些说法在现场记录中是否有明确判定依据？",
            "自然射线计数率的本底应优先用哪种方式确定：空载本底、纯煤稳定段、初始稳定段，还是随工作面推进动态更新？",
            "以 0.1 s 采样数据计算 ma10，并把它近似理解为 1 s 窗口后进行泊松置信区间分析，这种处理是否符合仪器计数逻辑？",
            "早期 data、新增 Excel 和 CSV 的计数率水平差异较大，我计划只做各自内部分析，不进行统一绝对阈值比较，这样是否合理？",
            "如果后续想研究前序支架影响，至少需要补充哪些现场记录？是否包括放煤口动作时间、支架编号、输送机速度、探测器安装位置和煤流到达时间？",
            "从本科论文角度看，您认为研究重点应放在“本底校正与统计判别”，还是“放煤阶段识别与可视化分析”？",
            "以上图形和表述中，哪些结论可以谨慎写入论文，哪些目前证据不足、不适合写成明确结论？",
        ]
    )

    add_heading(doc, "十一、下一步计划", 1)
    add_table(
        doc,
        ["任务", "拟开展内容", "预期产出"],
        [
            ["补充方法说明", "把 B(t)+R(t)、泊松置信区间、持续超限时间和阶段分割写成完整方法章节。", "论文方法章节初稿。"],
            ["完善图表体系", "统一图件编号、图注和指标口径，优先保留能回答研究问题的图。", "论文图件清单和可直接插入论文的图。"],
            ["加强结果讨论", "把结果分成可靠结论、谨慎推断和证据不足三类。", "降低过度解释风险。"],
            ["准备请教与补数", "根据老师意见决定是否补充空载本底、现场标签、支架动作和推进位置记录。", "形成下一轮数据需求清单。"],
        ],
        widths=[3.0, 7.2, 5.2],
    )

    add_body(
        doc,
        "以上是我目前的详细阶段性进展。由于我是本科生，专业理解和现场经验还不够成熟，恳请老师在方便时帮我指出研究方向、方法假设和现场解释中存在的问题。非常感谢老师的指导！",
        after=0,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
